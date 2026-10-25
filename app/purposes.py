"""Purpose mapping (purpose limitation readiness) module for Kavach.

Vocabulary and preset credit: teammate Nehar (SLM_Backend).
"""
import csv
import datetime
import json
import os
import re
from typing import Any, Dict, List, Optional

import yaml

from app.db import get_db

RULES_DIR = os.path.join(os.path.dirname(__file__), "rules")
PURPOSES_YAML_PATH = os.path.join(RULES_DIR, "purposes.yaml")
PRESET_YAML_PATH = os.path.join(RULES_DIR, "clinic_preset.yaml")

VALID_SEGMENTS = {"patient", "customer", "staff"}


def load_vocabulary() -> List[Dict[str, str]]:
    """Load the controlled vocabulary of business processing purposes."""
    with open(PURPOSES_YAML_PATH, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return data or []


def load_preset() -> Dict[str, Dict[str, List[str]]]:
    """Load the clinic preset mapping (segment, category) -> [purpose_ids]."""
    with open(PRESET_YAML_PATH, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return data or {}


def infer_segment(filename: str) -> str:
    """Infer customer segment from source filename.

    - patient -> patient
    - billing/customer -> customer
    - staff/employee -> staff
    - default -> customer
    """
    fn = (filename or "").lower()
    if "patient" in fn:
        return "patient"
    if "billing" in fn or "customer" in fn:
        return "customer"
    if "staff" in fn or "employee" in fn:
        return "staff"
    return "customer"


def _extract_source_headers(path: str, stype: str) -> Any:
    """Extract header mapping from source file on disk."""
    if not path or not os.path.exists(path):
        return {}
    if stype == "csv" or path.endswith(".csv"):
        try:
            with open(path, "r", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                headers = next(reader, [])
                return {i + 1: h.strip() for i, h in enumerate(headers)}
        except Exception:
            return {}
    elif stype == "xlsx" or path.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            res = {}
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    res[ws.title] = {i + 1: str(cell).strip() if cell is not None else "" for i, cell in enumerate(row)}
                    break
            wb.close()
            return res
        except Exception:
            return {}
    return {}


def get_questions(conn=None) -> Dict[str, Any]:
    """Generate open purpose questions for distinct fields (source, header, category).

    Skips fields with a confirmed active answer.
    Ordered by finding count descending.
    """
    if conn is None:
        conn = get_db()
    cur = conn.cursor()

    # 1. Load vocabulary and presets
    vocab = load_vocabulary()
    preset = load_preset()

    options = [
        {"id": item["id"], "label": item["label"], "label_key": item.get("label_key", item["id"])}
        for item in vocab
    ]
    options.append({"id": "unknown", "label": "I don't know", "label_key": "unknown"})

    # 2. Get sources
    cur.execute("SELECT id, filename, path, type FROM sources")
    sources = {s["id"]: dict(s) for s in cur.fetchall()}

    # 3. Extract header caches for sources
    source_headers = {sid: _extract_source_headers(s.get("path", ""), s.get("type", "")) for sid, s in sources.items()}

    # 4. Fetch all findings
    cur.execute("SELECT source_id, location, category, meta_json FROM findings")
    findings = cur.fetchall()

    field_counts: Dict[tuple, int] = {}
    for f in findings:
        sid = f["source_id"]
        s = sources.get(sid, {})
        loc = f["location"] or ""
        cat = f["category"]

        meta = {}
        if f["meta_json"]:
            try:
                meta = json.loads(f["meta_json"])
            except Exception:
                meta = {}

        header = meta.get("header")
        if not header:
            m_xlsx = re.match(r"^([^/]+)/R\d+/C(\d+)$", loc)
            m_csv = re.match(r"^R\d+/C(\d+)$", loc)
            if m_xlsx:
                sheet, col = m_xlsx.group(1), int(m_xlsx.group(2))
                header = source_headers.get(sid, {}).get(sheet, {}).get(col, "")
            elif m_csv:
                col = int(m_csv.group(1))
                header = source_headers.get(sid, {}).get(col, "")
            else:
                header = "(free text)"
        if not header:
            header = "(free text)"

        key = (sid, s.get("filename", ""), header, cat)
        field_counts[key] = field_counts.get(key, 0) + 1

    # 5. Fetch confirmed active answers
    cur.execute(
        "SELECT DISTINCT source_id, field_header, category FROM purpose_answers WHERE superseded_by IS NULL AND status = 'confirmed'"
    )
    confirmed_set = {(r["source_id"], r["field_header"], r["category"]) for r in cur.fetchall()}

    total_fields = len(field_counts)
    answered_count = sum(1 for (sid, _, h, cat) in field_counts if (sid, h, cat) in confirmed_set)

    questions = []
    for (sid, fn, header, category), count in field_counts.items():
        if (sid, header, category) in confirmed_set:
            continue

        segment = infer_segment(fn)

        # Preset suggestion (omit for aadhaar, unverified hint only)
        if category.lower() == "aadhaar":
            suggested_default = []
        else:
            seg_preset = preset.get(segment, {})
            hints = seg_preset.get(category)
            if hints is None and category == "health_term":
                hints = seg_preset.get("health", [])
            suggested_default = list(hints or [])

        questions.append({
            "source_id": sid,
            "source": fn,
            "source_filename": fn,
            "header": header,
            "field_header": header,
            "category": category,
            "segment": segment,
            "count": count,
            "finding_count": count,
            "suggested_default": suggested_default,
            "options": options,
        })

    # Sort descending by finding count, then by filename, header, category
    questions.sort(key=lambda q: (-q["finding_count"], q["source_filename"], q["field_header"], q["category"]))

    return {
        "questions": questions,
        "progress": {
            "answered": answered_count,
            "total": total_fields,
        },
    }


def record_answer(
    source_id: int,
    header: str,
    category: str,
    segment: str,
    purpose_ids: Optional[List[str]] = None,
    unknown: bool = False,
    answered_by: str = "user",
    conn=None,
) -> Dict[str, Any]:
    """Record an answer for a field purpose question.

    - Rejects invalid segments and unknown purpose IDs.
    - Supersedes prior active answers for the same (source_id, header, category, segment).
    - Emits ledger event 'purpose_answered' without raw personal data.
    """
    if segment not in VALID_SEGMENTS:
        raise ValueError(f"Invalid segment '{segment}'. Must be one of: {', '.join(sorted(VALID_SEGMENTS))}")

    if unknown:
        status = "unknown"
        pids_to_store = [None]
        purpose_ids_payload = []
    else:
        status = "confirmed"
        if not purpose_ids:
            raise ValueError("At least one purpose ID is required when unknown is False")
        if isinstance(purpose_ids, str):
            purpose_ids = [purpose_ids]
        vocab = load_vocabulary()
        valid_ids = {p["id"] for p in vocab}
        for pid in purpose_ids:
            if pid not in valid_ids:
                raise ValueError(f"Unknown purpose id: '{pid}'")
        pids_to_store = list(purpose_ids)
        purpose_ids_payload = list(purpose_ids)

    if conn is None:
        conn = get_db()
    cur = conn.cursor()

    now_ts = datetime.datetime.now(datetime.timezone.utc).isoformat()

    # Find prior active rows for this field & segment
    cur.execute(
        """SELECT id FROM purpose_answers
           WHERE source_id = ? AND field_header = ? AND category = ? AND segment = ? AND superseded_by IS NULL""",
        (source_id, header, category, segment),
    )
    prior_rows = cur.fetchall()
    prior_ids = [r["id"] for r in prior_rows]

    # Insert new row(s)
    new_ids = []
    for pid in pids_to_store:
        cur.execute(
            """INSERT INTO purpose_answers
               (source_id, field_header, category, segment, purpose_id, status, answered_by, answered_at, superseded_by)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, NULL)""",
            (source_id, header, category, segment, pid, status, answered_by, now_ts),
        )
        new_ids.append(cur.lastrowid)

    # Supersede prior rows
    if prior_ids and new_ids:
        first_new_id = new_ids[0]
        cur.execute(
            f"UPDATE purpose_answers SET superseded_by = ? WHERE id IN ({','.join('?' for _ in prior_ids)})",
            [first_new_id] + prior_ids,
        )

    # Get source filename
    cur.execute("SELECT filename FROM sources WHERE id = ?", (source_id,))
    s_row = cur.fetchone()
    filename = s_row["filename"] if s_row else ""

    # Ledger payload: NO raw values (source filename, header, category, segment, purpose ids, status only)
    payload = {
        "source_filename": filename,
        "field_header": header,
        "category": category,
        "segment": segment,
        "purpose_ids": purpose_ids_payload,
        "status": status,
    }
    from app.ledger import append_ledger
    append_ledger("purpose_answered", payload, conn)
    conn.commit()

    return {
        "inserted_ids": new_ids,
        "superseded_ids": prior_ids,
        "status": status,
    }


def get_history(source_id: int, header: str, category: str, conn=None) -> List[Dict[str, Any]]:
    """Retrieve full history of answers for a field (newest first)."""
    if conn is None:
        conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """SELECT id, source_id, field_header, category, segment, purpose_id, status, answered_by, answered_at, superseded_by
           FROM purpose_answers
           WHERE source_id = ? AND field_header = ? AND category = ?
           ORDER BY id DESC""",
        (source_id, header, category),
    )
    return [dict(r) for r in cur.fetchall()]
