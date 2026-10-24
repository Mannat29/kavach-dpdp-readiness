"""Erasure engine: build_plan, execute, and receipt generation."""
import csv
import hashlib
import html
import json
import os
import re
import shutil
import datetime
from typing import Dict, List, Optional, Tuple

import yaml

from app.db import get_db
from app.detect import scan_file, normalize_phone, hash_value
from app.ledger import append_ledger

_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_RULES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "rules")


def _load_retention_yaml() -> dict:
    path = os.path.join(_RULES_DIR, "retention.yaml")
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


# ---------------------------------------------------------------------------
# Build plan
# ---------------------------------------------------------------------------
def build_plan(request_id: int, conn=None) -> Dict:
    """Build erasure plan for a rights request. Saves plan_json + sha256 plan_hash."""
    if conn is None:
        conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT * FROM rights_requests WHERE id = ?", (request_id,))
    req = cur.fetchone()
    if not req:
        raise ValueError(f"Request {request_id} not found")

    person_id = req["person_id"]
    if not person_id:
        raise ValueError("Cannot build plan: person not resolved")

    retention = _load_retention_yaml()

    # Get all person records with their source info
    cur.execute("""
        SELECT pr.source_id, pr.record_key, pr.link_type,
               s.filename, s.path, s.type
        FROM person_records pr
        JOIN sources s ON pr.source_id = s.id
        WHERE pr.person_id = ?
    """, (person_id,))
    person_recs = cur.fetchall()

    plan = {"request_id": request_id, "person_id": person_id, "sources": {}}
    counts = {"ERASE": 0, "MASK": 0, "RETAIN": 0, "REVIEW": 0}

    for rec in person_recs:
        filename = rec["filename"]
        record_key = rec["record_key"]
        source_policy = retention.get("sources", {}).get(filename, {})

        # Get findings for this record
        cur.execute("""
            SELECT location, category, masked_sample, meta_json
            FROM findings
            WHERE source_id = ? AND record_key = ?
        """, (rec["source_id"], record_key))
        findings = cur.fetchall()

        source_plan = plan["sources"].setdefault(filename, {
            "path": rec["path"],
            "type": rec["type"],
            "records": {}
        })

        record_plan = {"fields": {}, "record_key": record_key}

        for f in findings:
            field_key = f["location"].split("/")[-1] if "/" in f["location"] else f["location"]
            # Try to get header from meta_json
            meta = json.loads(f["meta_json"]) if f["meta_json"] else {}
            header = meta.get("header", field_key)

            field_policy = _resolve_field_policy(source_policy, header, filename, f["category"])
            decision = field_policy["decision"]
            if decision == "REVIEW":
                counts["RETAIN"] += 1
            else:
                counts[decision] = counts.get(decision, 0) + 1

            record_plan["fields"][header] = {
                "decision": decision,
                "reason_code": field_policy["reason_code"],
                "citation_hint": field_policy["citation_hint"],
                "verify": field_policy.get("verify", False),
                "category": f["category"],
            }

        source_plan["records"][record_key] = record_plan

    plan["counts"] = counts
    plan_json = json.dumps(plan, sort_keys=True)
    plan_hash = hashlib.sha256(plan_json.encode("utf-8")).hexdigest()

    cur.execute("""
        UPDATE rights_requests SET plan_json = ?, plan_hash = ?, status = 'plan_ready'
        WHERE id = ?
    """, (plan_json, plan_hash, request_id))
    conn.commit()

    append_ledger({"event": "plan_built", "request_id": request_id, "plan_hash": plan_hash}, conn)
    return {"plan": plan, "plan_hash": plan_hash}


def _resolve_field_policy(source_policy: dict, header: str, filename: str, category: str) -> dict:
    """Resolve the decision for a field given source policy and header/category."""
    fields = source_policy.get("fields", {})

    # Try exact header match
    if header in fields:
        fp = fields[header]
        return {
            "decision": fp.get("decision", source_policy.get("default_decision", "RETAIN")),
            "reason_code": fp.get("reason_code", source_policy.get("reason_code", "UNKNOWN")),
            "citation_hint": fp.get("citation_hint", source_policy.get("citation_hint", "")),
            "verify": fp.get("verify", source_policy.get("verify", False)),
        }

    # Try category-based fallback for known types
    category_defaults = {
        "phone": {"decision": "MASK", "reason_code": "ID_MASK", "citation_hint": "identifier masked by default"},
        "email": {"decision": "ERASE", "reason_code": "ID_ERASE", "citation_hint": "identifier erased by default"},
        "aadhaar": {"decision": "ERASE", "reason_code": "ID_ERASE", "citation_hint": "identifier erased by default"},
        "name": {"decision": "ERASE", "reason_code": "ID_ERASE", "citation_hint": "identifier erased by default"},
        "dob": {"decision": "ERASE", "reason_code": "ID_ERASE", "citation_hint": "identifier erased by default"},
    }

    default_decision = source_policy.get("default_decision", "RETAIN")
    default_reason = source_policy.get("reason_code", "DEFAULT_RETAIN")
    default_citation = source_policy.get("citation_hint", "")
    default_verify = source_policy.get("verify", False)

    if category in category_defaults and default_decision != "RETAIN":
        cd = category_defaults[category]
        return {
            "decision": cd["decision"],
            "reason_code": cd["reason_code"],
            "citation_hint": cd["citation_hint"],
            "verify": default_verify,
        }

    return {
        "decision": default_decision,
        "reason_code": default_reason,
        "citation_hint": default_citation,
        "verify": default_verify,
    }


# ---------------------------------------------------------------------------
# Execute erasure
# ---------------------------------------------------------------------------
def execute(request_id: int, confirm_hash: str, conn=None) -> Dict:
    """
    Execute an erasure plan. Writes cleaned copies to data/out/<request_id>/.
    Originals are never modified. Verifies sha256 before and after.
    """
    if conn is None:
        conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT * FROM rights_requests WHERE id = ?", (request_id,))
    req = cur.fetchone()
    if not req:
        raise ValueError(f"Request {request_id} not found")

    if req["status"] != "approved":
        raise ValueError(f"Request {request_id} is not approved (status: {req['status']})")

    if req["plan_hash"] != confirm_hash:
        raise ValueError("confirm_hash does not match plan_hash — plan may have changed")

    plan = json.loads(req["plan_json"])
    person_id = req["person_id"]

    out_dir = os.path.join(_BASE_DIR, "data", "out", str(request_id))
    os.makedirs(out_dir, exist_ok=True)

    # Build person's raw identifiers IN MEMORY (never persisted)
    person_identifiers = _build_person_identifiers(person_id, conn)

    action_counts = {}
    retained_items = []
    residual_report = {}
    all_source_results = {}

    for filename, source_info in plan.get("sources", {}).items():
        src_path = source_info.get("path")
        if not src_path or not os.path.exists(src_path):
            continue

        file_ext = os.path.splitext(filename)[1].lower()

        # SHA256 of original
        original_hash = _sha256_file(src_path)

        # Build per-record decisions from plan
        record_decisions = {}  # record_key -> {field_header -> decision}
        for rk, rp in source_info.get("records", {}).items():
            record_decisions[rk] = {fh: fd["decision"] for fh, fd in rp["fields"].items()}
            for fh, fd in rp["fields"].items():
                d = fd["decision"]
                action_counts[d] = action_counts.get(d, 0) + 1
                if d in ("RETAIN",):
                    retained_items.append({
                        "source": filename,
                        "record_key": rk,
                        "field": fh,
                        "reason_code": fd["reason_code"],
                        "citation_hint": fd["citation_hint"],
                        "verify": fd.get("verify", False),
                    })

        out_path = os.path.join(out_dir, filename)

        if file_ext == ".xlsx":
            _process_xlsx(src_path, out_path, filename, record_decisions, plan)
        elif file_ext == ".csv":
            _process_csv(src_path, out_path, filename, record_decisions, plan)
        elif file_ext == ".txt":
            _process_txt(src_path, out_path, filename, record_decisions, plan, person_identifiers)

        # Verify original unchanged
        after_hash = _sha256_file(src_path)
        if original_hash != after_hash:
            raise RuntimeError(f"INTEGRITY ERROR: {src_path} was modified during execution!")

        # Residual check on cleaned file
        residuals = _residual_check(out_path, person_identifiers, record_decisions, source_info)
        residual_report[filename] = residuals

        all_source_results[filename] = {
            "original_sha256": original_hash,
            "cleaned_path": out_path,
        }

    # Also run scan_file on cleaned files and check for unexpected findings
    scan_residuals = {}
    for filename, sr in all_source_results.items():
        cp = sr["cleaned_path"]
        if os.path.exists(cp):
            scan_findings = scan_file(cp)
            scan_residuals[filename] = len(scan_findings)

    # Mark executed
    executed_at = datetime.datetime.utcnow().isoformat()
    cur.execute("""
        UPDATE rights_requests SET status = 'executed', executed_at = ?
        WHERE id = ?
    """, (executed_at, request_id))
    conn.commit()

    append_ledger({
        "event": "executed",
        "request_id": request_id,
        "plan_hash": confirm_hash,
        "executed_at": executed_at,
    }, conn)

    # Check if processors need notification
    processors_note = _get_processors_note(conn)

    # Write receipt
    receipt = {
        "request_id": request_id,
        "intent": req["intent"],
        "sender_masked": req["sender_masked"],
        "received_at": req["received_at"],
        "executed_at": executed_at,
        "plan_hash": confirm_hash,
        "action_counts": action_counts,
        "retained_items": retained_items,
        "residual_report": residual_report,
        "scan_residuals_per_file": scan_residuals,
        "processors_to_notify": processors_note,
        "note": "Backups and third-party copies are outside this tool's scope and must be handled separately.",
    }

    receipt_path = os.path.join(out_dir, "receipt.json")
    with open(receipt_path, encoding="utf-8", mode="w") as f:
        json.dump(receipt, f, indent=2, ensure_ascii=False)

    # Write HTML receipt
    _write_html_receipt(receipt, os.path.join(out_dir, "receipt.html"))

    return receipt


def _build_person_identifiers(person_id: int, conn) -> Dict:
    """Build raw identifiers IN MEMORY from person's original records. Never persist."""
    from app.detect import _PHONE_RE, _EMAIL_RE, _AADHAAR_RE
    from app.ingest import read_source
    from app.inventory import make_record_key

    cur = conn.cursor()

    identifiers = {"phones": set(), "emails": set(), "aadhaar_digits": set(), "name_tokens": set()}

    # Read from source files directly using inventory
    cur.execute("""
        SELECT DISTINCT s.path, s.filename
        FROM person_records pr
        JOIN sources s ON pr.source_id = s.id
        WHERE pr.person_id = ?
    """, (person_id,))
    sources = cur.fetchall()

    for src in sources:
        src_path = src["path"]
        if not src_path or not os.path.exists(src_path):
            continue

        # Get the record keys for this person from this source
        cur.execute("""
            SELECT record_key FROM person_records pr
            JOIN sources s ON pr.source_id = s.id
            WHERE pr.person_id = ? AND s.path = ?
        """, (person_id, src_path))
        person_record_keys = {r["record_key"] for r in cur.fetchall()}

        filename = os.path.basename(src_path)
        for location, text, meta in read_source(src_path):
            rk = make_record_key(filename, location)
            if rk not in person_record_keys:
                continue

            # Extract raw identifiers from text
            for m in _PHONE_RE.finditer(text):
                phone = normalize_phone(m.group())
                if phone:
                    identifiers["phones"].add(phone[-10:])

            for m in _EMAIL_RE.finditer(text):
                identifiers["emails"].add(m.group().lower())

            for m in _AADHAAR_RE.finditer(text):
                digits = m.group(1) + m.group(2) + m.group(3)
                identifiers["aadhaar_digits"].add(digits)

            # Name tokens from header-named fields
            header = meta.get("header", "").lower()
            if header in {"name", "patient_name"}:
                for word in text.split():
                    if len(word) >= 3:
                        identifiers["name_tokens"].add(word.lower())
            sender = meta.get("sender", "")
            if sender:
                for word in sender.split():
                    if len(word) >= 3:
                        identifiers["name_tokens"].add(word.lower())

    return identifiers


def _sha256_file(path: str) -> str:
    """Compute sha256 of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _process_xlsx(src_path: str, out_path: str, filename: str,
                  record_decisions: dict, plan: dict):
    """Process xlsx: blank erased cells, redact masked cells, keep retained cells."""
    from openpyxl import load_workbook
    wb = load_workbook(src_path)

    for ws in wb.worksheets:
        headers = []
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            if row_idx == 1:
                headers = [str(cell.value or "") for cell in row]
                continue

            # Build record_key for this row
            rk = f"{filename}:{ws.title}/R{row_idx}"

            if rk not in record_decisions:
                continue  # Not a person record, leave unchanged

            field_decisions = record_decisions[rk]

            for col_idx, cell in enumerate(row, start=1):
                header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                decision = field_decisions.get(header, "RETAIN")
                if decision == "ERASE":
                    cell.value = None
                elif decision == "MASK":
                    if cell.value:
                        cell.value = "[REDACTED]"

    wb.save(out_path)
    wb.close()


def _process_csv(src_path: str, out_path: str, filename: str,
                 record_decisions: dict, plan: dict):
    """Process csv: blank/mask fields per decision, skip erased rows only if ALL fields are ERASE."""
    rows = []
    with open(src_path, encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)

    if not rows:
        with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
            pass
        return

    headers = rows[0]
    out_rows = [headers]

    for row_idx, row in enumerate(rows[1:], start=2):
        rk = f"{filename}:R{row_idx}"

        if rk not in record_decisions:
            out_rows.append(row)
            continue

        field_decisions = record_decisions[rk]
        all_erase = all(d == "ERASE" for d in field_decisions.values()) and field_decisions

        if all_erase:
            continue  # Remove row

        new_row = list(row)
        for col_idx, header in enumerate(headers):
            decision = field_decisions.get(header, "RETAIN")
            if decision == "ERASE":
                new_row[col_idx] = ""
            elif decision == "MASK":
                if col_idx < len(new_row) and new_row[col_idx]:
                    new_row[col_idx] = "[REDACTED]"
        out_rows.append(new_row)

    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(out_rows)


def _process_txt(src_path: str, out_path: str, filename: str,
                 record_decisions: dict, plan: dict, person_identifiers: dict):
    """Process WhatsApp txt: erase sender's lines, mask identifiers in others' lines."""
    lines_to_erase = set()
    for rk, fields in record_decisions.items():
        if rk.startswith("line "):
            if fields and all(d == "ERASE" for d in fields.values()):
                try:
                    line_num = int(rk.split(" ")[1])
                    lines_to_erase.add(line_num)
                except (IndexError, ValueError):
                    pass

    out_lines = []
    with open(src_path, encoding="utf-8") as f:
        for line_num, raw_line in enumerate(f, start=1):
            line = raw_line.rstrip("\n\r")
            if line_num in lines_to_erase:
                continue  # Remove this line
            else:
                # Mask any of the person's identifiers appearing in this line
                masked = _mask_identifiers_in_text(line, person_identifiers)
                out_lines.append(masked)

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(out_lines))
        if out_lines:
            f.write("\n")


def _mask_identifiers_in_text(text: str, person_identifiers: dict) -> str:
    """Mask the person's identifiers in a text substring."""
    from app.detect import _PHONE_RE, _EMAIL_RE
    result = text

    # Mask phones that match person's phones
    def mask_phone_if_match(m):
        phone = normalize_phone(m.group())
        if phone and phone[-10:] in person_identifiers.get("phones", set()):
            return "[REDACTED]"
        return m.group()
    result = _PHONE_RE.sub(mask_phone_if_match, result)

    # Mask emails that match person's emails
    def mask_email_if_match(m):
        if m.group().lower() in person_identifiers.get("emails", set()):
            return "[REDACTED]"
        return m.group()
    result = _EMAIL_RE.sub(mask_email_if_match, result)

    return result


def _residual_check(out_path: str, person_identifiers: dict, record_decisions: dict, source_info: dict) -> dict:
    """
    Check cleaned file for residual PII.
    Returns report with expected (RETAIN decisions) and unexpected counts.
    """
    if not os.path.exists(out_path):
        return {"error": "cleaned file not found"}

    # Determine what was intentionally retained
    retained_records = set()
    for rk, fields in record_decisions.items():
        for fh, decision in fields.items():
            if decision in ("RETAIN",):
                retained_records.add(rk)

    with open(out_path, encoding="utf-8", errors="replace") as f:
        content = f.read()
    content_lower = content.lower()
    content_digits = re.sub(r"\D", "", content)

    report = {
        "phones": {"expected": 0, "unexpected": 0, "details": []},
        "emails": {"expected": 0, "unexpected": 0, "details": []},
        "aadhaar": {"expected": 0, "unexpected": 0, "details": []},
        "names": {"expected": 0, "unexpected": 0, "weak": True, "details": []},
    }

    # Check phones (digit-stripped)
    for phone in person_identifiers.get("phones", set()):
        stripped = re.sub(r"\D", "", phone)[-10:]
        if stripped in content_digits:
            is_expected = _is_expected_residual("phone", retained_records, source_info)
            if is_expected:
                report["phones"]["expected"] += 1
            else:
                report["phones"]["unexpected"] += 1
                report["phones"]["details"].append("phone found in cleaned file")

    # Check emails
    for email in person_identifiers.get("emails", set()):
        if email.lower() in content_lower:
            is_expected = _is_expected_residual("email", retained_records, source_info)
            if is_expected:
                report["emails"]["expected"] += 1
            else:
                report["emails"]["unexpected"] += 1
                report["emails"]["details"].append("email found in cleaned file")

    # Check Aadhaar digits
    for aadhaar in person_identifiers.get("aadhaar_digits", set()):
        stripped = re.sub(r"\D", "", aadhaar)
        if stripped in content_digits:
            is_expected = _is_expected_residual("aadhaar", retained_records, source_info)
            if is_expected:
                report["aadhaar"]["expected"] += 1
            else:
                report["aadhaar"]["unexpected"] += 1
                report["aadhaar"]["details"].append("aadhaar found in cleaned file")

    # Check name tokens (weak — common-name collision possible)
    for token in person_identifiers.get("name_tokens", set()):
        if len(token) >= 4 and re.search(rf"\b{re.escape(token)}\b", content_lower):
            is_expected = _is_expected_residual("name", retained_records, source_info)
            if is_expected:
                report["names"]["expected"] += 1
            else:
                report["names"]["unexpected"] += 1
                report["names"]["details"].append(f"name token '{token}' found (weak match)")

    return report


def _is_expected_residual(category: str, retained_records: set, source_info: dict) -> bool:
    """Determine if a residual is expected (due to RETAIN decisions)."""
    return bool(retained_records)


def _get_processors_note(conn) -> str:
    """Return a note about downstream processors that must also be notified."""
    return (
        "If personal data has been shared with any third-party processors "
        "(e.g. labs, cloud services, print vendors), notify them of this erasure request. "
        "This tool only cleans files in the local workspace."
    )


def _write_html_receipt(receipt: dict, out_path: str):
    """Write a simple HTML erasure receipt."""
    request_id = receipt.get("request_id", "?")
    intent = html.escape(receipt.get("intent", ""))
    sender_masked = html.escape(receipt.get("sender_masked", ""))
    received_at = html.escape(str(receipt.get("received_at", "")))
    executed_at = html.escape(str(receipt.get("executed_at", "")))
    plan_hash = html.escape(receipt.get("plan_hash", ""))
    action_counts = receipt.get("action_counts", {})
    retained_items = receipt.get("retained_items", [])
    note = html.escape(receipt.get("note", ""))
    processors_note = html.escape(receipt.get("processors_to_notify", ""))

    counts_html = "".join(
        f"<tr><td>{html.escape(k)}</td><td>{v}</td></tr>"
        for k, v in action_counts.items()
    )

    retained_html = ""
    if retained_items:
        retained_html = "<h3>Retained Items (require human review)</h3><table><tr><th>Source</th><th>Field</th><th>Reason</th><th>Citation</th></tr>"
        for item in retained_items:
            retained_html += (
                f"<tr><td>{html.escape(item.get('source',''))}</td>"
                f"<td>{html.escape(item.get('field',''))}</td>"
                f"<td>{html.escape(item.get('reason_code',''))}</td>"
                f"<td>{html.escape(item.get('citation_hint',''))}</td></tr>"
            )
        retained_html += "</table>"

    body = f"""<!DOCTYPE html>
<html lang="en">
<head><meta charset="utf-8"><title>Erasure Receipt #{request_id}</title>
<style>
body{{font-family:sans-serif;max-width:800px;margin:2em auto;padding:1em}}
table{{border-collapse:collapse;width:100%}}
th,td{{border:1px solid #ccc;padding:6px 10px;text-align:left}}
th{{background:#f0f0f0}}
.note{{background:#fff3cd;padding:1em;border-left:4px solid #ffc107;margin:1em 0}}
.processors{{background:#d1ecf1;padding:1em;border-left:4px solid #0c7cd5;margin:1em 0}}
</style>
</head>
<body>
<h1>Erasure Receipt — Request #{request_id}</h1>
<p><strong>Intent:</strong> {intent}</p>
<p><strong>Sender (masked):</strong> {sender_masked}</p>
<p><strong>Received:</strong> {received_at}</p>
<p><strong>Executed:</strong> {executed_at}</p>
<p><strong>Plan hash (sha256):</strong> <code>{plan_hash}</code></p>
<h3>Action Summary</h3>
<table><tr><th>Action</th><th>Count</th></tr>{counts_html}</table>
{retained_html}
<div class="note"><strong>Note:</strong> {note}</div>
<div class="processors"><strong>Processors:</strong> {processors_note}</div>
</body></html>
"""
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(body)
