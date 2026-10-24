"""Ingest module: read source files and yield (location, text, metadata) units."""
import csv
import logging
import os
import re

logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 50_000
ALLOWED_EXTENSIONS = {".xlsx", ".csv", ".txt"}


def read_source(path: str, max_rows: int = MAX_ROWS):
    """Yield (location, text, metadata) tuples from a file.

    location format:
      - xlsx: "SheetName/R{row}/C{col}"
      - csv:  "R{row}/C{col}"
      - txt (WhatsApp): "line {n}"

    metadata is a dict with optional keys: header, sender, row, col.
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    filename = os.path.basename(path)
    file_size = os.path.getsize(path)
    if file_size == 0:
        raise ValueError(f"File '{filename}' is empty (0 bytes).")
    if file_size > MAX_FILE_SIZE:
        raise ValueError(f"File '{filename}' exceeds maximum allowed size of 10 MB ({file_size:,} bytes).")

    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        yield from _read_xlsx(path, max_rows=max_rows)
    elif ext == ".csv":
        yield from _read_csv(path, max_rows=max_rows)
    elif ext == ".txt":
        yield from _read_txt(path, max_rows=max_rows)
    else:
        raise ValueError(f"Unsupported file type: '{ext}'. Only .xlsx, .csv, and .txt files are accepted.")


def _read_xlsx(path: str, max_rows: int = MAX_ROWS):
    filename = os.path.basename(path)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        logger.warning("Corrupted or malformed Excel file %s: %s", filename, e)
        raise ValueError(f"Corrupted or invalid Excel file '{filename}'.") from e

    try:
        total_rows = 0
        has_content = False
        for ws in wb.worksheets:
            headers = []
            for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
                total_rows += 1
                if total_rows > max_rows:
                    raise ValueError(f"File '{filename}' exceeds maximum limit of {max_rows:,} rows.")
                if row_idx == 1:
                    headers = [cell.value or "" for cell in row]
                    continue
                for col_idx, cell in enumerate(row, start=1):
                    val = cell.value
                    if val is None:
                        continue
                    text = str(val).strip()
                    if not text:
                        continue
                    has_content = True
                    header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                    location = f"{ws.title}/R{row_idx}/C{col_idx}"
                    yield location, text, {"header": str(header), "row": row_idx, "col": col_idx}
        if not has_content:
            raise ValueError(f"File '{filename}' contains no data rows.")
    finally:
        wb.close()


def _read_csv(path: str, max_rows: int = MAX_ROWS):
    filename = os.path.basename(path)
    try:
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            headers = []
            has_content = False
            for row_idx, row in enumerate(reader, start=1):
                if row_idx > max_rows:
                    raise ValueError(f"File '{filename}' exceeds maximum limit of {max_rows:,} rows.")
                if row_idx == 1:
                    headers = row
                    continue
                for col_idx, val in enumerate(row, start=1):
                    text = val.strip()
                    if not text:
                        continue
                    has_content = True
                    header = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
                    location = f"R{row_idx}/C{col_idx}"
                    yield location, text, {"header": str(header), "row": row_idx, "col": col_idx}
            if not has_content:
                raise ValueError(f"File '{filename}' contains no data rows.")
    except UnicodeDecodeError as e:
        logger.warning("Encoding error in CSV %s: %s", filename, e)
        raise ValueError(f"Malformed or invalid encoding in '{filename}'. File must be UTF-8 encoded.") from e
    except csv.Error as e:
        logger.warning("CSV format error in %s: %s", filename, e)
        raise ValueError(f"Malformed or corrupted CSV file '{filename}'.") from e


# WhatsApp Android export format: "dd/mm/yy, h:mm am - Sender: message"
_WA_RE = re.compile(
    r"^(\d{1,2}/\d{1,2}/\d{2,4},\s*\d{1,2}:\d{2}\s*[ap]m)\s*-\s*(.+?):\s*(.+)$",
    re.IGNORECASE,
)


def _read_txt(path: str, max_rows: int = MAX_ROWS):
    filename = os.path.basename(path)
    try:
        with open(path, "r", encoding="utf-8") as f:
            has_content = False
            for line_num, raw_line in enumerate(f, start=1):
                if line_num > max_rows:
                    raise ValueError(f"File '{filename}' exceeds maximum limit of {max_rows:,} lines.")
                line = raw_line.rstrip("\n\r")
                if not line:
                    continue
                has_content = True
                m = _WA_RE.match(line)
                if m:
                    sender = m.group(2).strip()
                    message = m.group(3).strip()
                    location = f"line {line_num}"
                    yield location, line, {"sender": sender, "message": message}
                else:
                    # System message or continuation
                    location = f"line {line_num}"
                    yield location, line, {}
            if not has_content:
                raise ValueError(f"File '{filename}' is empty or contains only whitespace.")
    except UnicodeDecodeError as e:
        logger.warning("Encoding error in text file %s: %s", filename, e)
        raise ValueError(f"Malformed or invalid encoding in '{filename}'. File must be UTF-8 encoded.") from e
